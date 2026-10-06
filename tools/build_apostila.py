import os,re,random
from pathlib import Path
from PIL import Image,ImageDraw,ImageFont
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path('.')
CAP=ROOT/'capitulos'; CAP.mkdir(exist_ok=True)
ILL=ROOT/'ilustracoes'; ILL.mkdir(exist_ok=True)
DOCX=ROOT/'Apostila_Munduruku_Leitura_Escrita_Edicao_Integral.docx'

new_titles={71:'Convenções ortográficas avançadas',72:'Nasalidade e grafemas especiais',73:'Apóstrofo e oclusiva glotal',74:'Estrutura silábica e leitura visual',75:'Paradigmas de posse',76:'Alomorfia dos marcadores pessoais',77:'Clíticos pessoais e prefixos relacionais',78:'Sintagma nominal em profundidade',79:'Sintagma verbal em profundidade',80:'Sintagma posposicional em profundidade',81:'Cisão de intransitivos',82:'Aspecto e indexação',83:'Reduplicação e interpretação',84:'Causativização e valência',85:'Nominalização e funções derivadas',86:'Incorporação nominal em leitura',87:'Nomes em função classificadora',88:'Posposições espaciais',89:'Posposições temporais e abstratas',90:'Coesão, anáfora e progressão textual',91:'Leitura de material documentado',92:'Dicionários, listas de palavras e corpus',93:'Variação ortográfica e textos históricos',94:'Tradução responsável e retroversão',95:'Escrita guiada por padrões confirmados',96:'Escrita autônoma',97:'Revisão e edição de texto',98:'Simulado avançado de leitura e escrita',99:'Projeto final de autonomia',100:'Gramática de consulta e plano pós-apostila'}
facts={71:'Compare a forma original da fonte com a forma normalizada e nunca altere a grafia sem registrar a mudança.',72:'A nasalidade é informação ortográfica relevante; trate vogais nasais e g̃ como sinais que precisam ser preservados.',73:'O apóstrofo representa a oclusiva glotal em descrições da ortografia e faz parte da palavra.',74:'Padrões V, CV, VC e CVC ajudam a decodificar visualmente, mas sílaba não é sinônimo de morfema.',75:'dao, a’õ e ba aparecem em paradigmas inalienáveis; kobe ilustra um padrão alienável.',76:'Marcadores pessoais apresentam alomorfia; compare paradigmas em vez de decorar uma forma única por pessoa.',77:'Gomes distingue clíticos pessoais de prefixos relacionais; nomes, verbos e posposições participam desses mecanismos.',78:'No sintagma nominal, identifique núcleo, possuidor e modificadores antes de traduzir.',79:'No sintagma verbal, identifique valência, aspecto, pessoa e argumentos expressos ou recuperáveis.',80:'Posposições têm complemento e exprimem relações espaciais, temporais e abstratas.',81:'Subclasses intransitivas podem apresentar padrões diferentes de indexação, sobretudo em interação com aspecto.',82:'O sufixo -m é descrito em usos imperfectivos; aspecto não deve ser reduzido a passado versus presente.',83:'Reduplicação pode copiar uma, duas ou três sílabas e pode indicar duração, pluralidade, iteração ou intensidade.',84:'Morfemas como mu-/muy- e do-/duju- aparecem em análises de causativização e mudança de valência.',85:'Nominalizadores como -ap e -at participam de construções em que bases verbais assumem funções nominais.',86:'Na incorporação, material nominal pode aparecer dentro do sintagma verbal e separar visualmente marcador e núcleo.',87:'Nomes em função classificadora precisam ser distinguidos do uso lexical independente.',88:'kay está ligado a direção; wi a origem; ase a relação superior, em usos documentados.',89:'wap pode ter valor temporal de anterioridade; em ocorre com duração ou frequência em usos descritos.',90:'Leitura longa exige rastrear referentes, tópico e informação recuperada pelo contexto.',91:'Texto documentado deve ser lido com registro de autor, edição, página e convenção ortográfica.',92:'Crofts e Sheffler oferecem dicionário escolar; a Munduruku Word List é muito mais extensa e tem finalidade diferente.',93:'Materiais históricos podem usar convenções diferentes; preserve a forma original e justifique qualquer normalização.',94:'Traduza unidades de sentido e estrutura, não uma sequência de equivalentes palavra por palavra.',95:'Escreva a partir de padrões confirmados e mude um elemento por vez.',96:'Na escrita autônoma, reduza o português como etapa intermediária e planeje com imagens, esquemas e formas confirmadas.',97:'Revise separadamente grafia, segmentação, morfologia, sintaxe, coesão e escolha lexical.',98:'O simulado avançado exige análise justificada, consulta responsável e indicação explícita de incerteza.',99:'O projeto final integra leitura documental, análise morfológica, resumo, escrita e rastreamento de fontes.',100:'Depois do curso, avance por corpus: textos, dicionários, gramáticas e um léxico pessoal com fonte e status de confirmação.'}
lex=['kobe','dao','a’õ','ba','uk’a','op','abikbikap','parawa','bio','ag̃oka','wamõat','wãtaxipi','muketero','xik','’at','dakat','nomuwã','jeorok','jepidowat','ajok','cu','õn','ẽn','wuyju','oceju','eyju','pug̃','xepxep','ebapug','ebadipdip','soat','ade','xere','wara’at','abu','ajo','poce','poma','podi','pẽn','peburu','pebit','kay','wi','ase','pibun','wap','be','eju','buxim','em','g̃u','acã','tak','koap','waram','xipan']

def slug(s):
    tr=str.maketrans('áàãâéêíóôõúç','aaaaeeiooouc')
    return re.sub(r'[^a-z0-9]+','-',s.lower().translate(tr)).strip('-')

for n,t in new_titles.items():
    p=CAP/f'{n:02d}-{slug(t)}.md'
    text=f'''# Capítulo {n} — {t}\n\n## Meta\n\nAprender a ler e escrever com foco em {t.lower()}, sem treino obrigatório de fala.\n\n## Explicação central\n\n{facts[n]}\n\nA regra de segurança desta unidade é separar três coisas: o que a forma mostra, o que o contexto sugere e o que a fonte confirma. Quando uma dessas camadas faltar, marque a análise como provisória.\n\n## Método\n\n1. Observe a forma inteira.\n2. Marque sinais gráficos e partes já confirmadas.\n3. Identifique a função provável na construção.\n4. Faça uma hipótese contextual.\n5. Confirme em gramática, dicionário ou texto-fonte.\n6. Releia a unidade inteira.\n\n## Exercícios\n\n- Explique o mecanismo com suas próprias palavras.\n- Escolha cinco formas confirmadas e registre classe, função e fonte.\n- Anote duas informações que ainda dependem de contexto.\n- Refaça a análise no dia seguinte sem consultar a resposta anterior.\n\n## Critério de domínio\n\nAvance quando conseguir justificar a análise e distinguir dado confirmado de hipótese.\n\n### Fontes-base\n\nCrofts; Gomes; Picanço; Crofts & Sheffler; estudos posteriores listados na bibliografia da edição integral.\n'''
    p.write_text(text,encoding='utf-8')

levels=[(1,0,10,'Fundamentos'),(2,11,20,'Estrutura intermediária'),(3,21,30,'Leitura e escrita avançadas'),(4,31,40,'Morfologia avançada'),(5,41,50,'Autonomia orientada'),(6,51,60,'Leitura aplicada'),(7,61,70,'Automatização gramatical'),(8,71,80,'Ortografia e sintagmas em profundidade'),(9,81,90,'Morfossintaxe avançada e discurso'),(10,91,100,'Autonomia documental e produção escrita')]

def chapter_files():
    items=[]
    for p in CAP.glob('*.md'):
        m=re.match(r'^(\d+)-',p.name)
        if m: items.append((int(m.group(1)),p))
    return sorted(items)

items=chapter_files(); titles={}
for n,p in items:
    first=p.read_text(encoding='utf-8').splitlines()[0].lstrip('# ').strip()
    titles[n]=re.sub(r'^Capítulo\s+\d+\s*[—-]\s*','',first)
idx=['# Índice geral — edição integral','']
for lv,a,b,name in levels:
    idx += [f'## Nível {lv} — {name}','']
    for n in range(a,b+1): idx.append(f'{n:02d}. {titles.get(n,new_titles.get(n,""))}')
    idx.append('')
idx += ['## Arquivos finais','','- `Apostila_Munduruku_Leitura_Escrita_Edicao_Integral.docx`','- `Apostila_Munduruku_Leitura_Escrita_Edicao_Integral.pdf`']
(CAP/'indice-geral.md').write_text('\n'.join(idx),encoding='utf-8')

fontpath='/usr/share/fonts/truetype/noto/NotoSerif-Regular.ttf'
if not Path(fontpath).exists(): fontpath='/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf'
for name,lab in [('edicao-integral-leitura','ver → segmentar → classificar → inferir → confirmar → reler'),('edicao-integral-morfologia','pessoa • raiz • aspecto • partícula'),('edicao-integral-espaco','origem ← referente → direção'),('edicao-integral-progresso','palavra → sintagma → oração → parágrafo → texto')]:
    im=Image.new('RGB',(1400,700),(247,241,224));d=ImageDraw.Draw(im);f=ImageFont.truetype(fontpath,46)
    d.rectangle((60,60,1340,640),outline=(34,76,55),width=5);d.text((110,115),'MUNDURUKÚ — LEITURA E ESCRITA',font=f,fill=(20,45,34))
    f2=ImageFont.truetype(fontpath,34);d.text((110,330),lab,font=f2,fill=(75,80,70));d.line((110,250,1270,250),fill=(176,137,73),width=4)
    im.save(ILL/f'{name}.png',optimize=True)

doc=Document(); sec=doc.sections[0];sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.72);sec.right_margin=Inches(.72)
for st in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
    s=doc.styles[st];s.font.name='Noto Serif';s._element.rPr.rFonts.set(qn('w:ascii'),'Noto Serif');s._element.rPr.rFonts.set(qn('w:hAnsi'),'Noto Serif')
doc.styles['Normal'].font.size=Pt(10.2);doc.styles['Title'].font.size=Pt(28);doc.styles['Title'].font.color.rgb=RGBColor(34,76,55)
for st in ['Heading 1','Heading 2','Heading 3']:doc.styles[st].font.color.rgb=RGBColor(34,76,55)

def page_no(p):
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run();a=OxmlElement('w:fldChar');a.set(qn('w:fldCharType'),'begin');b=OxmlElement('w:instrText');b.text=' PAGE ';c=OxmlElement('w:fldChar');c.set(qn('w:fldCharType'),'end');r._r.extend([a,b,c])
foot=sec.footer.paragraphs[0];foot.text='Mundurukú — leitura e escrita • edição integral';foot.runs[0].font.size=Pt(8);page_no(sec.footer.add_paragraph())

def title(k,t):
    p=doc.add_paragraph();r=p.add_run(k.upper());r.bold=True;r.font.size=Pt(9);r.font.color.rgb=RGBColor(70,95,72)
    p=doc.add_paragraph();p.style='Title';p.add_run(t)
def body(s):
    p=doc.add_paragraph(s);p.paragraph_format.space_after=Pt(6);p.paragraph_format.line_spacing=1.05

def lines(prompts,n=2):
    for i,x in enumerate(prompts,1):
        p=doc.add_paragraph();r=p.add_run(f'{i}. {x}');r.bold=True;r.font.size=Pt(9.5)
        for _ in range(n): doc.add_paragraph('________________________________________________________________________________')

title('Curso autodidata','MUNDURUKÚ')
body('Leitura e escrita — do zero à autonomia avançada')
body('Sem foco em conversação. Conteúdo organizado por fontes, análise, revisão ativa e escrita responsável.')
doc.add_picture(str(ILL/'edicao-integral-leitura.png'),width=Inches(6.4));doc.add_page_break()
title('Princípio','Como estudar')
body('Primeiro veja a estrutura. Depois confirme o sentido. O português diminui progressivamente, mas a fonte nunca desaparece: quando uma forma não estiver confirmada, deixe a lacuna explícita.')
body('A repetição desta edição é pedagógica: cada retorno muda a operação — reconhecer, reconstruir, explicar, analisar, escrever e revisar.')
doc.add_picture(str(ILL/'edicao-integral-progresso.png'),width=Inches(6.4));doc.add_page_break()

random.seed(20261006)
for n,p in items:
    raw=p.read_text(encoding='utf-8')
    titletext=titles[n]
    title(f'Capítulo {n:02d}',titletext)
    for line in raw.splitlines()[1:]:
        s=line.strip()
        if not s or s=='---':continue
        if s.startswith('### '):doc.add_heading(s[4:],level=3)
        elif s.startswith('## '):doc.add_heading(s[3:],level=2)
        elif s.startswith('# '):doc.add_heading(s[2:],level=1)
        elif s.startswith('- '):doc.add_paragraph(s[2:],style='List Bullet')
        elif re.match(r'^\d+\. ',s):doc.add_paragraph(re.sub(r'^\d+\. ','',s),style='List Number')
        else: body(s.replace('**','').replace('`',''))
    doc.add_page_break()
    for sheet in range(1,9):
        names=['Leitura guiada','Desmontagem morfológica','Recuperação ativa','Consulta de fontes','Produção controlada','Revisão de erros','Revisão espaçada','Teste de domínio']
        title(f'Capítulo {n:02d} • prática {sheet}',names[sheet-1])
        chosen=random.sample(lex,8)
        body('Formas para recuperação: '+' • '.join(chosen))
        prompts=['Observe as formas e registre apenas o que consegue afirmar sem adivinhar.','Separe grafia, possível estrutura e sentido contextual em três colunas mentais.','Escolha duas formas e diga que informação ainda depende de contexto.','Relacione o tema deste capítulo a um conteúdo estudado anteriormente.','Escreva qual fonte você consultaria para confirmar sua hipótese.','Registre um erro real e reescreva a resposta depois da correção.']
        if sheet in (5,8):prompts+=['Planeje uma produção escrita sem criar primeiro uma frase portuguesa completa.','Marque o que está confirmado, provável e incerto.']
        lines(prompts,1 if sheet not in (5,8) else 2)
        doc.add_page_break()

for lv,a,b,name in levels:
    for k in range(1,21):
        title(f'Revisão do nível {lv} • ficha {k:02d}',name)
        body(f'Revisão intercalada dos capítulos {a:02d}–{b:02d}. Responda de memória antes de consultar.')
        body('Formas de recuperação: '+' • '.join(random.sample(lex,10)))
        lines(['Classifique as formas sem traduzir imediatamente.','Reconstrua cinco ideias centrais deste nível.','Explique uma relação entre morfologia e sintaxe.','Identifique uma dúvida que exige fonte.','Reescreva uma resposta antiga melhorando a justificativa.','Planeje a próxima revisão espaçada.'],1)
        doc.add_page_break()

title('Apêndice','Léxico-base para exercícios')
for i in range(0,len(lex),3):body(' • '.join(lex[i:i+3]))
doc.add_page_break()
title('Apêndice','Fontes principais')
for s in ['Crofts, Marjorie. Gramática Mundurukú.','Gomes, Dioney Moreira. Estudo morfológico e sintático da língua Mundurukú (Tupí). UnB, 2006.','Picanço, Gessiane Lobato. Mundurukú: Phonetics, Phonology, Synchrony, Diachrony. UBC, 2005.','Crofts, Marjorie; Sheffler, Margaret. Dicionário Bilíngüe em Português e Mundurukú. 1981.','Gomes, Dioney M. Postpositions in Munduruku (Tupi): Formal and Functional Features. 2019.','Nóbrega de Abreu, Natali; Picanço, Gessiane Lobato. Alomorfia dos prefixos pessoais de posse nominal da língua Mundurukú. 2022.','Borges, Renan do Socorro dos Santos; Lopes, Jorge Domingues. Para uma crítica lexicográfica das microestruturas da Munduruku Word List. 2018.','Intercontinental Dictionary Series — Mundurukú, baseado em Crofts & Sheffler 1981.']:body(s)

doc.save(DOCX)
print('capitulos',len(items),'docx',DOCX)
